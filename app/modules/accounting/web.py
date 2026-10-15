from __future__ import annotations

import csv
import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.async_db import get_async_session
from app.core.models import Client, Supplier
from app.core.permissions import PERMISSION_OPERATIONS_READ, PERMISSION_OPERATIONS_WRITE
from app.core.rate_limit import limiter
from app.modules.accounting.scf_service import SCFService
from app.modules.accounting.service import transactions_context
from app.modules.production.web import edit_production_notes  # noqa: F401 (compat)
from app.modules.purchases.service import PurchaseService
from app.modules.sales.service import SalesService
from app.modules.payments.service import new_payment_context
from app.infra.printing import COMPANY_INFO, PRINT_LAYOUT, build_print_payload, generate_invoice_pdf
from app.web.deps import (
    flash,
    get_current_user,
    login_redirect,
    require_permission,
    template_context,
    templates,
)

router = APIRouter()


@router.get("/accounting/scf", name="accounting_scf_dashboard")
@limiter.limit("20/minute")
async def accounting_dashboard(
    request: Request,
    db: AsyncSession = Depends(get_async_session),
):
    user = get_current_user(request)
    if not user:
        from fastapi.responses import RedirectResponse

        return RedirectResponse("/login", status_code=303)

    balance = await SCFService.get_balance_generale(db=db)
    tcr = await SCFService.get_tcr(balance=balance, db=db)
    bilan = await SCFService.get_bilan(balance=balance, db=db)

    return templates.TemplateResponse(
        "accounting_scf.html",
        template_context(request, title="Comptabilité SCF Algérie", balance=balance, tcr=tcr, bilan=bilan),
    )


@router.get("/accounting/balance/export", name="accounting_balance_export")
@limiter.limit("10/minute")
async def export_balance(
    request: Request,
    db: AsyncSession = Depends(get_async_session),
):
    user = get_current_user(request)
    if not user:
        from fastapi.responses import RedirectResponse

        return RedirectResponse("/login", status_code=303)

    balance = await SCFService.get_balance_generale(db=db)

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(
        [
            "Compte SCF",
            "Intitulé du Compte",
            "Débit (DA)",
            "Crédit (DA)",
            "Solde Débiteurs (DA)",
            "Solde Créditeur (DA)",
        ]
    )

    total_deb = 0.0
    total_cred = 0.0
    total_solde_deb = 0.0
    total_solde_cred = 0.0

    for acc in balance:
        writer.writerow(
            [
                acc["code"],
                acc["label"],
                f"{acc['debit']:.2f}",
                f"{acc['credit']:.2f}",
                f"{acc['solde_debiteur']:.2f}",
                f"{acc['solde_crediteur']:.2f}",
            ]
        )
        total_deb += acc["debit"]
        total_cred += acc["credit"]
        total_solde_deb += acc["solde_debiteur"]
        total_solde_cred += acc["solde_crediteur"]

    writer.writerow([])
    writer.writerow(
        [
            "TOTAL",
            "TOTAL GÉNÉRAL BALANCE",
            f"{total_deb:.2f}",
            f"{total_cred:.2f}",
            f"{total_solde_deb:.2f}",
            f"{total_solde_cred:.2f}",
        ]
    )

    output.seek(0)
    bom = "\ufeff"
    return StreamingResponse(
        iter([bom + output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=balance_generale_scf_{date.today().isoformat()}.csv"},
    )


@router.get("/accounting/scf/export-excel", name="accounting_scf_export_excel")
@router.get("/accounting/export-excel", name="accounting_export_excel")
@limiter.limit("10/minute")
async def export_scf_excel(
    request: Request,
    db: AsyncSession = Depends(get_async_session),
):
    user = get_current_user(request)
    if not user:
        from fastapi.responses import RedirectResponse

        return RedirectResponse("/login", status_code=303)

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    balance = await SCFService.get_balance_generale(db=db)
    tcr = await SCFService.get_tcr(balance=balance, db=db)
    bilan = await SCFService.get_bilan(balance=balance, db=db)

    wb = openpyxl.Workbook()

    # Style presets
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    section_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="000000")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Sheet: Balance Générale
    ws_bal = wb.active
    ws_bal.title = "Balance Générale"
    ws_bal.append(
        [
            "Compte SCF",
            "Intitulé du Compte",
            "Débit (DA)",
            "Crédit (DA)",
            "Solde Débiteur (DA)",
            "Solde Créditeur (DA)",
        ]
    )

    for col in range(1, 7):
        cell = ws_bal.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    tot_deb, tot_cred, tot_sdeb, tot_scred = 0.0, 0.0, 0.0, 0.0
    for row_idx, acc in enumerate(balance, start=2):
        ws_bal.append(
            [
                acc["code"],
                acc["label"],
                acc["debit"],
                acc["credit"],
                acc["solde_debiteur"],
                acc["solde_crediteur"],
            ]
        )
        for col in range(1, 7):
            c = ws_bal.cell(row=row_idx, column=col)
            c.border = thin_border
            if col in (3, 4, 5, 6):
                c.number_format = '#,##0.00 "DA"'
                c.alignment = align_right

        tot_deb += acc["debit"]
        tot_cred += acc["credit"]
        tot_sdeb += acc["solde_debiteur"]
        tot_scred += acc["solde_crediteur"]

    tot_row = len(balance) + 2
    ws_bal.append(["TOTAL", "TOTAL GÉNÉRAL BALANCE", tot_deb, tot_cred, tot_sdeb, tot_scred])
    for col in range(1, 7):
        c = ws_bal.cell(row=tot_row, column=col)
        c.font = total_font
        c.fill = section_fill
        c.border = thin_border
        if col in (3, 4, 5, 6):
            c.number_format = '#,##0.00 "DA"'
            c.alignment = align_right

    # 2. Sheet: Bilan (Actif vs Passif)
    ws_bilan = wb.create_sheet(title="Bilan Comptable")
    ws_bilan.append(["POSTES DE L'ACTIF", "MONTANT (DA)", "", "POSTES DU PASSIF", "MONTANT (DA)"])
    for col in (1, 2, 4, 5):
        cell = ws_bilan.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    actif_rows = [
        ("Stocks de matières & produits", bilan["actif"]["stocks"]),
        ("Créances Clients", bilan["actif"]["creances_clients"]),
        ("Trésorerie & Caisse", bilan["actif"]["tresorerie_caisse"]),
    ]
    passif_rows = [
        ("Capitaux Propres", bilan["passif"]["capitaux_propres"]),
        ("Résultat Net de l'Exercice", bilan["passif"]["resultat_exercice"]),
        ("Dettes Fournisseurs", bilan["passif"]["dettes_fournisseurs"]),
    ]

    for i in range(max(len(actif_rows), len(passif_rows))):
        a_lbl, a_val = actif_rows[i] if i < len(actif_rows) else ("", "")
        p_lbl, p_val = passif_rows[i] if i < len(passif_rows) else ("", "")
        ws_bilan.append([a_lbl, a_val, "", p_lbl, p_val])
        row_num = i + 2
        for col, val in ((1, a_lbl), (2, a_val), (4, p_lbl), (5, p_val)):
            c = ws_bilan.cell(row=row_num, column=col)
            c.border = thin_border
            if col in (2, 5) and isinstance(val, (int, float)):
                c.number_format = '#,##0.00 "DA"'
                c.alignment = align_right

    ws_bilan.append(["TOTAL ACTIF", bilan["actif"]["total"], "", "TOTAL PASSIF", bilan["passif"]["total"]])
    tot_bilan_row = max(len(actif_rows), len(passif_rows)) + 2
    for col in (1, 2, 4, 5):
        c = ws_bilan.cell(row=tot_bilan_row, column=col)
        c.font = total_font
        c.fill = section_fill
        c.border = thin_border
        if col in (2, 5):
            c.number_format = '#,##0.00 "DA"'
            c.alignment = align_right

    # 3. Sheet: TCR (Compte de Résultat)
    ws_tcr = wb.create_sheet(title="Compte de Résultat (TCR)")
    ws_tcr.append(["RUBRIQUE TCR (SCF)", "MONTANT (DA)"])
    for col in (1, 2):
        cell = ws_tcr.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    tcr_items = [
        ("Chiffre d'Affaires (Ventes 700 / 707)", tcr["chiffre_affaires"]),
        ("Achats consommés (600 / 607)", tcr["achats_consommes"]),
        ("Marge Brute", tcr["marge_brute"]),
        ("Autres charges externes d'exploitation (625 / 629)", tcr["autres_charges_externes"]),
        ("Valeur Ajoutée d'Exploitation", tcr["valeur_ajoutee"]),
        ("Résultat d'Exploitation", tcr["resultat_exploitation"]),
        ("RÉSULTAT NET DE L'EXERCICE", tcr["resultat_net"]),
    ]

    for idx, (lbl, val) in enumerate(tcr_items, start=2):
        ws_tcr.append([lbl, val])
        c1 = ws_tcr.cell(row=idx, column=1)
        c2 = ws_tcr.cell(row=idx, column=2)
        c1.border = thin_border
        c2.border = thin_border
        c2.number_format = '#,##0.00 "DA"'
        c2.alignment = align_right
        if "RÉSULTAT" in lbl or "Marge Brute" in lbl or "Valeur Ajoutée" in lbl:
            c1.font = total_font
            c2.font = total_font
            c1.fill = section_fill
            c2.fill = section_fill

    # Auto-adjust column widths
    for sheet in (ws_bal, ws_bilan, ws_tcr):
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 14)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    filename = f"etats_financiers_scf_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _print_not_found_response(message: str = "Bon introuvable.") -> HTMLResponse:
    return HTMLResponse(
        f"""<!doctype html>
<html lang="fr">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"></head>
<body style="margin:0;font-family:Arial,sans-serif;background:#f8fafc;color:#111827;display:grid;place-items:center;min-height:100vh;">
  <main style="max-width:520px;padding:24px;text-align:center;">
    <h1 style="font-size:1.25rem;margin:0 0 8px;">{message}</h1>
    <p style="margin:0;color:#64748b;">Le document demandé n'existe plus ou n'est pas disponible.</p>
  </main>
</body>
</html>""",
        status_code=404,
    )


@router.get("/operations", name="operations")
@router.get("/transactions", name="transactions")
async def operations_page(request: Request, db: AsyncSession = Depends(get_async_session)):
    denied = require_permission(request, PERMISSION_OPERATIONS_READ)
    if denied:
        return denied

    fmt = request.query_params.get("format", "").lower()
    if fmt in ("csv", "xlsx"):
        large_args = dict(request.query_params)
        large_args["page_size"] = 1000000
        context = await transactions_context(
            filter_type=str(request.query_params.get("type", "all") or "all"),
            filter_name=str(request.query_params.get("name", "") or ""),
            filter_date=str(request.query_params.get("date", "") or ""),
            filter_operation=str(request.query_params.get("operation", "") or ""),
            args=large_args,
            path=request.url.path,
            db=db,
        )
        data = context["transactions"]

        if fmt == "csv":
            output = io.StringIO()
            fieldnames = [
                "tx_type",
                "tx_date",
                "partner_name",
                "designation",
                "quantity",
                "unit",
                "unit_price",
                "total",
                "paid",
                "due",
            ]
            writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";", extrasaction="ignore")
            writer.writerow({f: f.upper() for f in fieldnames})
            for row in data:
                writer.writerow(row)
            output.seek(0)
            filename = f"transactions_{date.today().isoformat()}.csv"
            return StreamingResponse(
                iter(["\ufeff" + output.getvalue()]),
                media_type="text/csv; charset=utf-8",
                headers={"Content-Disposition": f'attachment; filename="{filename}"'},
            )

        elif fmt == "xlsx":
            import openpyxl
            from openpyxl.styles import Font

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Transactions"

            headers = [
                "TYPE",
                "DATE",
                "TIERS",
                "DÉSIGNATION",
                "QUANTITÉ",
                "UNITÉ",
                "PRIX UNITAIRE",
                "TOTAL",
                "PAYÉ",
                "DU",
            ]
            ws.append(headers)

            fieldnames = [
                "tx_type",
                "tx_date",
                "partner_name",
                "designation",
                "quantity",
                "unit",
                "unit_price",
                "total",
                "paid",
                "due",
            ]
            for row in data:
                ws.append([str(row.get(f) or "") if row.get(f) is not None else "" for f in fieldnames])

            ws.freeze_panes = "A2"
            bold_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = bold_font

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

            filename = f"transactions_{date.today().isoformat()}.xlsx"
            out_buf = io.BytesIO()
            wb.save(out_buf)
            out_buf.seek(0)
            return StreamingResponse(
                out_buf,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f'attachment; filename="{filename}"'},
            )

    context = await transactions_context(
        filter_type=str(request.query_params.get("type", "all") or "all"),
        filter_name=str(request.query_params.get("name", "") or ""),
        filter_date=str(request.query_params.get("date", "") or ""),
        filter_operation=str(request.query_params.get("operation", "") or ""),
        args=request.query_params,
        path=request.url.path,
        db=db,
    )
    return templates.TemplateResponse("transactions.html", template_context(request, **context))


@router.get("/operations/new", name="new_operation")
async def new_operation_page(request: Request, db: AsyncSession = Depends(get_async_session)):
    denied = require_permission(request, PERMISSION_OPERATIONS_WRITE)
    if denied:
        return denied

    p_ctx = await PurchaseService(db).purchase_form_context()
    s_ctx = await SalesService(db).sale_form_context()
    pay_ctx = await new_payment_context(db=db)

    context = {}
    context.update(p_ctx)
    context.update(s_ctx)
    context.update(pay_ctx)

    clients_res = await db.execute(select(Client).order_by(Client.name))
    context["clients"] = [c.model_dump() for c in clients_res.scalars().all()]

    suppliers_res = await db.execute(select(Supplier).order_by(Supplier.name))
    context["suppliers"] = [s.model_dump() for s in suppliers_res.scalars().all()]
    context["mode"] = request.query_params.get("mode", "achat")

    return templates.TemplateResponse("operation_new.html", template_context(request, **context))


@router.get("/print/{doc_type}/{item_id}", name="print_document")
async def print_document_page(
    request: Request, doc_type: str, item_id: int, db: AsyncSession = Depends(get_async_session)
):
    user = get_current_user(request)
    if not user:
        return login_redirect()
    payload = await build_print_payload(doc_type, item_id, db=db)
    if not payload:
        return _print_not_found_response("Document introuvable pour impression.")
    printed_at = datetime.now()
    payload = {
        **payload,
        "printed_date": printed_at.strftime("%Y-%m-%d"),
        "printed_time": printed_at.strftime("%H:%M"),
    }
    if request.query_params.get("format") == "pdf":
        pdf_buf = generate_invoice_pdf(payload, user["username"])
        if pdf_buf:
            headers = {"Content-Disposition": f'attachment; filename="{payload["number"]}.pdf"'}
            return StreamingResponse(pdf_buf, media_type="application/pdf", headers=headers)
        flash(request, "Génération PDF indisponible. Affichage HTML utilisé à la place.", "warning")
    return templates.TemplateResponse(
        "print_document.html",
        template_context(
            request,
            doc=payload,
            company=payload.get("company") or COMPANY_INFO,
            printed_by=user["username"],
            print_layout=PRINT_LAYOUT,
        ),
    )

